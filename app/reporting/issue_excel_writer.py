"""Standalone Excel-Report ausschließlich aus ``ValidationIssue``-Liste.

Phase-1.5-Standardpfad: Reporting konsumiert *nur* die Issue-Objekte und
die Summary – kein Legacy-Validator-State, keine Side-Effects. Layout ist
bewusst schlicht (drei Sheets: Issues, Summary, Manifest), damit Tests
deterministisch sind und das umfangreiche Legacy-Layout in
``excel_report_writer.write_excel_report`` weiterhin als optionaler
Vollreport zur Verfügung steht.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Optional

from app.models import ValidationIssue, ValidationSummary
from app.reporting.manifest import RunManifest

_ISSUE_COLUMNS = [
    "rule_id",
    "rule_level",
    "rule_type",
    "template",
    "row",
    "field_code",
    "field_label",
    "severity",
    "value",
    "message",
    "explanation",
    "dpm_reference",
    "timestamp",
]


def write_issue_excel(
    issues: Iterable[ValidationIssue],
    output_path: str | Path,
    summary: Optional[ValidationSummary] = None,
    manifest: Optional[RunManifest] = None,
) -> str:
    """Schreibt einen Standalone-Excel-Report aus den Issue-Objekten."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "openpyxl wird für den Excel-Report benötigt."
        ) from exc

    wb = openpyxl.Workbook()
    ws_issues = wb.active
    ws_issues.title = "Issues"
    ws_issues.append(_ISSUE_COLUMNS)
    for it in issues:
        row = it.to_dict()
        ws_issues.append([row.get(c, "") for c in _ISSUE_COLUMNS])

    ws_sum = wb.create_sheet("Summary")
    if summary is not None:
        ws_sum.append(["metric", "value"])
        for key, val in summary.to_dict().items():
            ws_sum.append([key, str(val)])

    ws_man = wb.create_sheet("Manifest")
    if manifest is not None:
        ws_man.append(["key", "value"])
        for k, v in manifest.to_dict().items():
            ws_man.append([k, str(v)])

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)
