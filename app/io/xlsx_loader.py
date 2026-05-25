"""XLSX-Eingaben laden (SRB Annex I Format)."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

import openpyxl
import pandas as pd

from app.io.template_registry import normalize_sheet_name
from app.models import InputBatch, TemplateData
from app.normalization.field_resolver import resolve_dataframe_columns
from app.normalization.headers import extract_headers, find_header_row
from app.normalization.values import normalize_dataframe_strings


def load_xlsx(
    filepath: str | Path,
    batch: Optional[InputBatch] = None,
    field_structure: Any = None,
) -> InputBatch:
    """Lädt ein MBDT-XLSX-Workbook (SRB Annex I Format)."""
    batch = batch or InputBatch(source_type="xlsx", source_path=str(filepath))
    wb = openpyxl.load_workbook(filepath, data_only=True)

    for sheet_name in wb.sheetnames:
        normalized = normalize_sheet_name(sheet_name)
        if not normalized:
            continue
        ws = wb[sheet_name]
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            continue
        header_idx = find_header_row(all_rows)
        if header_idx is None:
            continue
        headers = extract_headers(all_rows, header_idx)
        if not headers:
            continue

        data_rows = []
        for row in all_rows[header_idx + 1:]:
            if any(cell is not None for cell in row):
                row_list = list(row)
                if len(row_list) < len(headers):
                    row_list += [None] * (len(headers) - len(row_list))
                else:
                    row_list = row_list[: len(headers)]
                data_rows.append(row_list)

        if data_rows:
            df = pd.DataFrame(data_rows, columns=headers)
            df = normalize_dataframe_strings(df)
            if field_structure is not None:
                df, _diag = resolve_dataframe_columns(df, normalized, field_structure)
            batch.add(TemplateData(
                template_id=normalized,
                key=normalized,
                df=df,
                variant=None,
                source_name=sheet_name,
            ))

    return batch
